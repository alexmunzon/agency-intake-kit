"""PR 4 readers: encoding, delimiter, header row, trailing totals, and lineage."""

import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from agency_schema.formats import parse_date_loose
from intake.readers import RawTable
from intake.readers.csv import read_csv
from intake.readers.sniff import decode, find_header_row, is_total_row, sniff_delimiter
from intake.readers.xlsx import cell_text, read_xlsx

DROP = Path(__file__).parents[3] / "fixtures" / "agency-a" / "drop"
MANIFEST = json.loads((DROP / "manifest.json").read_text())["files"]
HEX64 = re.compile(r"^[0-9a-f]{64}$")


def read(name: str, sheet: str | None = None) -> RawTable:
    if sheet is None:
        return read_csv(DROP / name, source="s", run_id="r1", mapping_version="unmapped")
    return read_xlsx(DROP / name, sheets={sheet: "s"}, run_id="r1", mapping_version="unmapped")[0]


@pytest.mark.parametrize("entry", MANIFEST, ids=lambda e: f"{e['file_name']}:{e['sheet']}")
def test_every_fixture_file_reads_with_manifest_row_count_and_lineage(entry: dict) -> None:
    table = read(entry["file_name"], entry["sheet"])
    assert table.rows == entry["rows"]
    lineage = table.frame["lineage"].struct.unnest()
    assert set(lineage["source_file"]) == {entry["file_name"]}
    assert lineage["sheet"].to_list() == [entry["sheet"]] * table.rows
    rows = lineage["row_number"].to_list()
    assert rows == list(range(table.header_row + 1, table.header_row + 1 + table.rows))
    assert all(HEX64.match(h) for h in lineage["raw_hash"])
    assert set(lineage["run_id"]) == {"r1"} and set(lineage["mapping_version"]) == {"unmapped"}
    assert all(
        dtype == pl.String for name, dtype in table.frame.schema.items() if name != "lineage"
    )


def test_raw_hash_is_stable_across_two_reads() -> None:
    first, second = read("crm_export.csv"), read("crm_export.csv")
    assert first.frame["lineage"].struct.field("raw_hash").to_list() == (
        second.frame["lineage"].struct.field("raw_hash").to_list()
    )


def test_bom_followed_by_latin1_body_decodes_as_latin1() -> None:
    data = (DROP / "crm_export.csv").read_bytes()
    assert data.startswith(b"\xef\xbb\xbf")
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8-sig")  # a reader that trusts the marker fails here
    text, encoding = decode(data)
    assert encoding == "latin-1" and not text.startswith("﻿")
    assert "ñ" in text or "é" in text


def test_crm_reads_accents_and_notes_ing_001() -> None:
    table = read("crm_export.csv")
    assert table.frame.columns[0] == "Client ID"  # no BOM glued to the first header
    notes = table.frame["Notes"].drop_nulls().str.join(" ").item()
    assert "�" not in notes and re.search("[ñé]", notes)
    assert [r.rule_id for r in table.exceptions] == ["ING-001"]
    assert "latin-1" in table.exceptions[0].message


def test_enrollment_semicolon_detected_and_utf8_is_quiet() -> None:
    table = read("enrollment_export.csv")
    assert (table.delimiter, table.encoding) == (";", "utf-8")
    assert "Birth Dt (mm/dd/yy)" in table.frame.columns
    assert table.exceptions == ()


@pytest.mark.parametrize(
    ("name", "first_header"),
    [("commissions_bluepeak.xlsx", "Stmt Period"), ("commissions_northwind_health.xlsx", "Line")],
)
def test_statement_header_on_row_3_past_merged_title_and_total_dropped(
    name: str, first_header: str
) -> None:
    table = read(name, "Statement")
    assert table.header_row == 3 and table.frame.columns[0] == first_header
    assert table.dropped_rows == 1 and table.total_row_count == table.rows
    assert table.frame["lineage"].struct.field("row_number")[0] == 4
    assert [r.rule_id for r in table.exceptions] == ["ING-002", "ING-003"]
    assert all(r.severity == "INFO" and r.row_number is None for r in table.exceptions)


def test_date_cells_become_iso_and_text_cells_stay_text() -> None:
    northwind = read("commissions_northwind_health.xlsx", "Statement")
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2}", northwind.frame["DOB"][0])
    assert cell_text(datetime(2026, 7, 15)) == "2026-07-15"
    assert cell_text(datetime(2026, 7, 15, 9, 30)) == "2026-07-15T09:30:00"
    assert cell_text(date(2026, 7, 15)) == "2026-07-15"
    assert cell_text("45901") == "45901" and cell_text(None) is None


def test_small_files_title_row_total_row_and_low_confidence_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "t.csv"
    path.write_text(
        "Agency report,,\nName,Policy,Amount\nAnn,P-1,1.00\nBo,P-2,2.00\nTotal,2,3.00\n,,\n"
    )
    table = read_csv(path, source="t", run_id="r", mapping_version="m")
    assert (table.header_row, table.rows, table.dropped_rows, table.total_row_count) == (2, 2, 2, 2)
    assert table.frame["lineage"].struct.field("row_number").to_list() == [3, 4]
    assert sniff_delimiter("a;b\nc;d\n").delimiter == ";"
    assert not sniff_delimiter("one\ntwo\n").confident
    assert find_header_row([["Title", None, None], ["a", "b", "c"], ["1", "2", "3"]]) == 1


def test_xlsx_without_sheet_list_reads_every_sheet(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.active.append(["Name", "When"])
    book.active.append(["Ann", datetime(2026, 1, 31)])
    book.create_sheet("Other").append(["Only", "Header"])
    book.save(tmp_path / "b.xlsx")
    tables = read_xlsx(tmp_path / "b.xlsx", sheets=None, run_id="r", mapping_version="m")
    assert [(t.source, t.rows) for t in tables] == [("b", 1), ("b", 0)]
    assert tables[0].frame["When"].to_list() == ["2026-01-31"]


@pytest.mark.parametrize(
    "first",
    [
        "Total",
        "TOTALS",
        "  totals ",
        "Total:",
        "TOTAL (3 periods)",
        "Grand Total",
        "grand  total",
        "Subtotal",
        "Sub Total",
        "sub-total",
    ],
)
def test_total_row_first_cell_spellings(first: str) -> None:
    assert is_total_row([first, "3", "157.50"])


@pytest.mark.parametrize("first", ["Totally Covered LLC", "Ann Total", "", None, "P-1"])
def test_names_that_only_contain_total_are_not_total_rows(first: str | None) -> None:
    assert not is_total_row([first, "3", "157.50"])


def _statement(tmp_path: Path, tail: list[list[str | None]]) -> RawTable:
    book = openpyxl.Workbook()
    rows = [["Acme Statement"], ["Line", "Member", "Amount"]]
    rows += [["1", "M-1", "50.00"], ["2", "M-2", "50.00"], ["3", "M-3", "57.50"]]
    for row in rows + tail:
        book.active.append(row)
    book.save(tmp_path / "s.xlsx")
    return read_xlsx(tmp_path / "s.xlsx", sheets=None, run_id="r", mapping_version="m")[0]


@pytest.mark.parametrize(
    "tail",
    [
        [["Total", "3", "157.50"], ["Generated by CarrierPortal"]],
        [["Totals", "3", "157.50"], ["Generated by CarrierPortal"]],
        [["Total:", "3", "157.50"], [None, None, None], ["Page 1 of 1"]],
        [["Subtotal", "3", "157.50"], ["Grand Total", "3", "157.50"], ["Report run 2026-10-04"]],
    ],
)
def test_total_block_followed_by_footer_lines_is_dropped(
    tmp_path: Path, tail: list[list[str | None]]
) -> None:
    table = _statement(tmp_path, tail)
    assert (table.rows, table.dropped_rows, table.total_row_count) == (3, len(tail), 3)
    assert table.frame["Member"].to_list() == ["M-1", "M-2", "M-3"]
    assert any(r.rule_id == "ING-003" for r in table.exceptions)


def test_a_footer_line_without_a_total_row_stays_as_data(tmp_path: Path) -> None:
    table = _statement(tmp_path, [["Generated by CarrierPortal"], [None, None, None]])
    assert (table.rows, table.dropped_rows, table.total_row_count) == (4, 1, None)


def test_every_enrollment_effective_date_parses() -> None:
    table = read("enrollment_export.csv")
    values = table.frame["effective"].drop_nulls().to_list()
    assert len(values) == table.rows == 1847
    assert all(parse_date_loose(v) is not None for v in values), "compact yyyymmdd must parse"
