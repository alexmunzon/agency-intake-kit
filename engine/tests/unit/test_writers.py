import csv
import hashlib
import io
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from functools import cache
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from synth_agency_data.cli import app
from synth_agency_data.writers.common import EXCEL_EPOCH, SERIAL_RANGE

FIXTURES = Path(__file__).parents[3] / "fixtures"
FLAGS = {
    "agency-a": [],
    "agency-a-truncated": ["--truncate-crm", "2574", "--no-canonical"],
    "agency-a-ssn": ["--add-ssn-column", "--no-canonical"],
}
CRM_ROWS = 2680  # 2,600 policies + 34 duplicate rows + 46 clients with no policy row


def _digest(root: Path) -> dict[str, str]:
    files = sorted(p for p in root.rglob("*") if p.is_file())
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}


@pytest.mark.parametrize("name", sorted(FLAGS))
def test_fixture_regenerates_byte_for_byte(name: str, tmp_path: Path) -> None:
    args = ["generate", "--seed", "42", "--out", str(tmp_path), *FLAGS[name]]
    result = CliRunner().invoke(app, args)
    assert result.exit_code == 0, result.output
    assert _digest(tmp_path) == _digest(FIXTURES / name), f"regenerate fixtures/{name}"


@cache
def crm(name: str) -> tuple[bytes, list[dict[str, str]]]:
    raw = (FIXTURES / name / "drop" / "crm_export.csv").read_bytes()
    return raw, list(csv.DictReader(io.StringIO(raw[3:].decode("latin-1"))))


@cache
def sheet(name: str, file: str, title: str) -> list[list[Any]]:
    wb = load_workbook(FIXTURES / name / "drop" / file)
    return [list(r) for r in wb[title].iter_rows(values_only=True)]


@cache
def truth(name: str) -> list[dict[str, Any]]:
    defects: list[dict[str, Any]] = json.loads((FIXTURES / name / "ground_truth.json").read_text())[
        "defects"
    ]
    return defects


def manifest(name: str) -> list[dict[str, Any]]:
    files: list[dict[str, Any]] = json.loads(
        (FIXTURES / name / "drop" / "manifest.json").read_text()
    )["files"]
    return files


def test_crm_is_latin1_with_a_bom_serial_dobs_and_mixed_dates() -> None:
    raw, rows = crm("agency-a")
    assert raw.startswith(b"\xef\xbb\xbf")
    with pytest.raises(UnicodeDecodeError):
        raw[3:].decode("utf-8")  # the body really is latin-1 ("espa\xf1ol")
    assert len(rows) == CRM_ROWS
    serials = [r["Mbr DOB"] for r in rows if r["Mbr DOB"].isdigit()]
    assert serials and all(SERIAL_RANGE[0] <= int(s) <= SERIAL_RANGE[1] for s in serials)
    shapes = Counter(
        re.sub(r"\d", "9", r["Eff Date"]) for r in rows if r["Eff Date"]
    )  # 99/99/9999, 9999-99-99, 99/99/99
    assert set(shapes) == {"99/99/9999", "9999-99-99", "99/99/99"}
    assert {"Active", "ACTIVE", "active"} <= {r["Status"] for r in rows}


def test_enrollment_is_semicolon_with_two_digit_birth_years_after_1930() -> None:
    text = (FIXTURES / "agency-a" / "drop" / "enrollment_export.csv").read_text("utf-8")
    rows = list(csv.DictReader(io.StringIO(text), delimiter=";"))
    assert text.splitlines()[0].count(";") == 9 and "Birth Dt (mm/dd/yy)" in rows[0]
    dobs = {r["Policy #"]: r["Mbr DOB"] for r in crm("agency-a")[1] if r["Policy #"]}
    for r in rows:
        if not r["Birth Dt (mm/dd/yy)"]:
            continue
        mm, dd, yy = (int(x) for x in r["Birth Dt (mm/dd/yy)"].split("/"))
        year = 1900 + yy if yy >= 30 else 2000 + yy  # the 1930 to 2029 pivot
        born = date(year, mm, dd)
        assert born.year >= 1930
        crm_dob = dobs[r["policy_number"]]
        expected = (
            EXCEL_EPOCH + timedelta(int(crm_dob))
            if crm_dob.isdigit()
            else date(int(crm_dob[6:]), int(crm_dob[:2]), int(crm_dob[3:5]))
        )
        assert born == expected


def test_statements_have_a_merged_title_header_on_row_3_and_a_total_row() -> None:
    for entry in manifest("agency-a"):
        if not entry["source"].startswith("statement_"):
            continue
        file = entry["file_name"]
        ws = load_workbook(FIXTURES / "agency-a" / "drop" / file)["Statement"]
        assert [str(r) for r in ws.merged_cells.ranges] == ["A1:J1"]
        rows = sheet("agency-a", file, "Statement")
        assert "Commission Statement" in rows[0][0] and rows[1][0].startswith("Agency A")
        header, data, total = rows[2], rows[3:-1], rows[-1]
        assert all(isinstance(h, str) for h in header) and len(header) == 10
        assert total[0] == "Total" and total[1] == str(len(data)) == str(entry["rows"])
        amounts = [r[-1] for r in data]
        assert all(re.fullmatch(r"-?\d+\.\d\d", a) for a in amounts), file
        cents = sum(int(a.replace(".", "")) for a in amounts)
        assert total[-1] == f"{cents // 100}.{cents % 100:02d}"


def test_roster_has_comma_license_lists_and_one_rts_row_per_combination() -> None:
    agents = sheet("agency-a", "agent_roster.xlsx", "Agents")
    rts = sheet("agency-a", "agent_roster.xlsx", "RTS")
    assert "SSN" not in agents[0] and "Licensed States" in agents[0]
    assert any(", " in r[3] for r in agents[1:])
    assert len({tuple(r[:5]) for r in rts[1:]}) == len(rts) - 1
    assert isinstance(rts[1][7], datetime) and rts[1][7].year >= 2022


@pytest.mark.parametrize("name", sorted(FLAGS))
def test_manifest_counts_match_the_full_files(name: str) -> None:
    for entry in manifest(name):
        file = entry["file_name"]
        if file == "crm_export.csv":
            assert entry["rows"] == CRM_ROWS  # the truncated fixture still says the full count
            got = len(crm(name)[1])
            assert got == (2574 if name == "agency-a-truncated" else CRM_ROWS)
        elif file.endswith(".csv"):
            text = (FIXTURES / name / "drop" / file).read_text("utf-8")
            assert len(text.splitlines()) - 1 == entry["rows"]
        else:
            rows = sheet(name, file, entry["sheet"])
            skip = 4 if entry["source"].startswith("statement_") else 1  # title rows, total row
            assert len(rows) - skip == entry["rows"]


@pytest.mark.parametrize("name", sorted(FLAGS))
def test_every_defect_points_at_a_row_that_holds_its_record(name: str) -> None:
    rows = crm(name)[1]
    cut = 2574 if name == "agency-a-truncated" else None
    for d in truth(name):
        key, row = d["record_key"], d["source_row"]
        if d["defect_type"] in ("truncated_file", "ssn_column"):
            continue
        if d["defect_type"] == "statement_total_variance":  # a TIE-005 total spans every file
            assert (d["source_file"], d["sheet"], row) == (None, None, None), d
            continue
        if d["source_file"] == "crm_export.csv":
            if cut is not None and row is None:
                continue  # its row was cut off; the check below proves no row was dropped early
            assert row is not None and (cut is None or row <= cut + 1), d
            cells = rows[row - 2]
            column = "Policy #" if "policy_id" in key else "Client ID"
            assert cells[column] == next(iter(key.values())), d
            if d["defect_type"] == "pii_in_notes":
                assert cells["Notes"] == d["injected_values"]["to"]
        else:
            cells = sheet(name, d["source_file"], d["sheet"])[row - 1]
            header = sheet(name, d["source_file"], d["sheet"])[2]
            got = {h: cells[i] for i, h in enumerate(header)}
            assert (
                d["source_file"]
                == "commissions_" + key["carrier"].lower().replace(" ", "_") + ".xlsx"
            )
            assert key["statement_period"] in got.values() and str(key["line_no"]) in got.values()
    if cut is not None:  # exactly the defects whose full-file row was cut lost their row
        full = [d for d in truth("agency-a")]
        for d, f in zip(truth(name), full, strict=False):
            if f["source_file"] == "crm_export.csv":
                assert (d["source_row"] is None) == (f["source_row"] > cut + 1), d


def test_truncated_fixture_records_the_cut_in_ground_truth() -> None:
    (lost,) = [d for d in truth("agency-a-truncated") if d["defect_type"] == "truncated_file"]
    assert lost["expected_rule_ids"] == ["CMP-001"] and lost["source_row"] == 2575
    assert lost["injected_values"] == {"rows_expected": CRM_ROWS, "rows_received": 2574}
    assert (
        (FIXTURES / "agency-a-truncated" / "drop" / "crm_export.csv").read_bytes().endswith(b"\r\n")
    )


def test_ssn_fixture_values_are_fake_and_in_ground_truth() -> None:
    (ssn,) = [d for d in truth("agency-a-ssn") if d["defect_type"] == "ssn_column"]
    assert ssn["expected_rule_ids"] == ["SSN-001"] and ssn["source_row"] == 1
    agents = sheet("agency-a-ssn", "agent_roster.xlsx", "Agents")
    column = agents[0].index("SSN")
    values = [r[column] for r in agents[1:]]
    assert values == ssn["injected_values"]["values"] and len(set(values)) == 25
    assert all(re.fullmatch(r"000-\d\d-\d{4}", v) for v in values)  # area 000 is never issued
    others = [d for d in truth("agency-a") if d["defect_type"] != "ssn_column"]
    assert [d for d in truth("agency-a-ssn") if d["defect_type"] != "ssn_column"] == others


def test_pii_notes_are_one_percent_scored_and_obviously_fake() -> None:
    pii = [d for d in truth("agency-a") if d["defect_type"] == "pii_in_notes"]
    assert len(pii) == 26 and all(
        d["scored"] and d["expected_rule_ids"] == ["PII-001"] for d in pii
    )
    assert len({d["record_key"]["policy_id"] for d in pii}) == 26
    keys = Counter(json.dumps(d["record_key"]) for d in truth("agency-a"))
    assert all(keys[json.dumps(d["record_key"])] == 1 for d in pii)
    for d in pii:
        sentence = d["injected_values"]["to"]
        assert not re.search(r"\d{3}-?\d{2}-?\d{4}", sentence.replace("555-01", "")), sentence


def test_exact_duplicate_rows_are_byte_identical_in_the_crm_file() -> None:
    """Issue 54: a DUP-001 copy repeats its original's bytes (same date style, same casing).

    DUP-003 copies differ in content, so they must still differ in the file.
    """
    raw, rows = crm("agency-a")
    lines = raw[3:].split(b"\r\n")[1:-1]  # drop the BOM, the header, and the final empty piece
    assert len(lines) == len(rows)  # no cell holds a line break, so line i is row i
    planted = {"exact_duplicate_row": 0, "duplicate_policy_id": 0}
    for d in truth("agency-a"):
        if d["defect_type"] not in planted:
            continue
        planted[d["defect_type"]] += 1
        pid = d["record_key"]["policy_id"]
        first, copy = [lines[i] for i, r in enumerate(rows) if r["Policy #"] == pid]
        if d["defect_type"] == "exact_duplicate_row":
            assert first == copy, pid
        else:
            assert first != copy, pid
    assert planted == {"exact_duplicate_row": 26, "duplicate_policy_id": 8}
