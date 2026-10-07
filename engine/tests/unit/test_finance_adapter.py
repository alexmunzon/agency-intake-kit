import hashlib

import pytest
from openpyxl import Workbook
from test_revenue_ledger import mapping

from intake.finance_adapter import StatementAdapter, import_statement
from intake.readers import raw_hash
from intake.revenue import classify_statement


def config(**kwargs):
    return StatementAdapter(
        adapter_version="layout-v1",
        agency_id="synthetic",
        carrier="Synthetic Carrier",
        period="2026-09",
        statement_id="s",
        revision_id="r1",
        run_id="run-synthetic",
        amount_column="Value",
        category_column="Label",
        transaction_column="Kind",
        **kwargs,
    )


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_signed_unknown_and_provenance(tmp_path, kind):
    path = tmp_path / f"synthetic.{kind}"
    cells = [
        ["Value", "Label", "Kind"],
        ["10.00", "Renewal", "Pay"],
        ["-2.50", "Unfamiliar", "Reverse"],
    ]
    if kind == "csv":
        path.write_text("\n".join(",".join(row) for row in cells))
    else:
        book = Workbook()
        for row in cells:
            book.active.append(row)
        book.save(path)
    package = import_statement(path, config(format=kind))
    review = classify_statement(package, mapping())
    assert str(review.statement_total) == "7.50"
    assert sum(review.category_totals.values()) == review.statement_total
    assert str(review.category_totals["unclassified"]) == "-2.50"
    assert review.rows[1].raw_category_label == "Unfamiliar"
    assert review.rows[1].classification_method == "unresolved"
    assert package.content_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert package.rows[1].lineage.raw_hash == raw_hash(cells[2])
    assert package.rows[1].lineage.row_number == 3
    assert package.rows[1].lineage.mapping_version == "layout-v1"
    assert import_statement(path, config(format=kind)) == package


@pytest.mark.parametrize(
    "body",
    [
        "Value,Label,Kind\n1.001,Renewal,Pay",
        "Value,Label,Kind\n1,Renewal,Pay,extra",
        "Value,Label,Label\n1,Renewal,Pay",
        "Value,Label,Kind\nNaN,Renewal,Pay",
    ],
)
def test_malformed_csv_refused(tmp_path, body):
    path = tmp_path / "bad.csv"
    path.write_text(body)
    with pytest.raises(ValueError):
        import_statement(path, config(format="csv"))


def test_explicit_header_footer_and_formula_refusal(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    for row in [
        ["Synthetic title"],
        ["Value", "Label", "Kind"],
        ["2.00", "Renewal", "Pay"],
        ["total", None, None],
    ]:
        book.active.append(row)
    book.save(path)
    assert (
        len(import_statement(path, config(format="xlsx", header_row=2, last_data_row=3)).rows) == 1
    )
    book.active["A3"] = "=1+1"
    book.save(path)
    with pytest.raises(ValueError, match="formula"):
        import_statement(path, config(format="xlsx", header_row=2, last_data_row=3))


def test_explicitly_excluded_xlsx_formula_footer_is_not_imported(tmp_path):
    path = tmp_path / "statement.xlsx"
    book = Workbook()
    for row in [
        ["Value", "Label", "Kind"],
        ["2.00", "Renewal", "Pay"],
        ["=SUM(A2:A2)", "Total", None],
    ]:
        book.active.append(row)
    book.save(path)
    package = import_statement(path, config(format="xlsx", last_data_row=2))
    assert len(package.rows) == 1
    assert package.rows[0].raw_amount == "2.00"
    with pytest.raises(ValueError, match="formula"):
        import_statement(path, config(format="xlsx"))


def test_stale_source_pin_refused(tmp_path):
    path = tmp_path / "statement.csv"
    path.write_text("Value,Label,Kind\n1.00,Renewal,Pay\n")
    cfg = config(format="csv", source_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    import_statement(path, cfg)
    path.write_text("Value,Label,Kind\n2.00,Renewal,Pay\n")
    with pytest.raises(ValueError, match="pinned"):
        import_statement(path, cfg)


def test_source_decimal_notation_is_preserved(tmp_path):
    path = tmp_path / "statement.csv"
    path.write_text("Value,Label,Kind\n+1e2,Renewal,Pay\n")
    package = import_statement(path, config(format="csv"))
    assert package.rows[0].raw_amount == "+1e2"
    assert package.rows[0].amount == 100
