"""Read xlsx sheets into raw frames: values only, every cell text except real date cells.

openpyxl returns a date-typed cell as a datetime; that cell becomes an ISO string (a bare
date when the time is midnight). Every other cell stays the text it was, so "45901" stored as
text is still "45901". Numbers stored as numbers become their plain string form.
"""

from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from intake.readers import RawTable, table_from_rows


def cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def read_xlsx(
    path: Path, *, sheets: dict[str, str] | None, run_id: str, mapping_version: str
) -> list[RawTable]:
    """One RawTable per sheet. `sheets` maps sheet name to source name; None reads them all."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        wanted = sheets if sheets is not None else {s: path.stem for s in book.sheetnames}
        tables = []
        for name, source in wanted.items():
            if name not in book.sheetnames:
                continue
            rows = [[cell_text(v) for v in row] for row in book[name].iter_rows(values_only=True)]
            tables.append(
                table_from_rows(
                    rows,
                    source=source,
                    source_file=path.name,
                    sheet=name,
                    run_id=run_id,
                    mapping_version=mapping_version,
                )
            )
        return tables
    finally:
        book.close()
