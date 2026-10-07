"""Explicit CSV/XLSX layouts to the existing statement evidence contract."""

import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Annotated, Literal

from openpyxl import load_workbook
from pydantic import Field

from agency_schema.lineage import Lineage, NonEmpty, Sha256
from intake.readers import raw_hash
from intake.revenue import ExactMoney, FinanceInput, RevenueInputRow, StatementPackage


class StatementAdapter(FinanceInput):
    source_sha256: Sha256 | None = None
    adapter_version: NonEmpty
    agency_id: NonEmpty
    carrier: NonEmpty
    period: str
    statement_id: NonEmpty
    revision_id: NonEmpty
    run_id: NonEmpty
    format: Literal["csv", "xlsx"]
    amount_column: NonEmpty
    category_column: NonEmpty
    transaction_column: NonEmpty
    row_id_column: NonEmpty | None = None
    header_row: Annotated[int, Field(strict=True, ge=1)] = 1
    last_data_row: Annotated[int | None, Field(strict=True, ge=1)] = None
    sheet: NonEmpty | None = None
    delimiter: Annotated[str, Field(min_length=1, max_length=1)] = ","
    encoding: str = "utf-8-sig"
    control_total: ExactMoney | None = None


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def import_statement(path: Path, config: StatementAdapter) -> StatementPackage:
    """Read bytes once, retain raw cells, and refuse the entire malformed input.

    XLSX formulas in the selected header/data region are refused, including cached
    values. No sign/category inference. Explicit bounds can exclude title/footer rows.
    """
    data = path.read_bytes()
    if (
        config.source_sha256 is not None
        and hashlib.sha256(data).hexdigest() != config.source_sha256
    ):
        raise ValueError("source bytes do not match the pinned SHA-256")
    sheet = None
    if config.format == "csv":
        if config.sheet is not None:
            raise ValueError("CSV cannot select a sheet")
        table = list(
            csv.reader(
                io.StringIO(data.decode(config.encoding)), delimiter=config.delimiter, strict=True
            )
        )
    else:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            if config.sheet is None:
                if len(book.sheetnames) != 1:
                    raise ValueError("select an explicit XLSX sheet")
                sheet = book.sheetnames[0]
            else:
                sheet = config.sheet
            if sheet not in book.sheetnames:
                raise ValueError("configured sheet is missing")
            table = []
            for number, cells in enumerate(book[sheet].iter_rows(), 1):
                in_import_region = number >= config.header_row and (
                    config.last_data_row is None or number <= config.last_data_row
                )
                if in_import_region and any(cell.data_type == "f" for cell in cells):
                    raise ValueError("formula cells are unsupported")
                table.append([_cell(cell.value) for cell in cells])
        finally:
            book.close()
    if len(table) < config.header_row:
        raise ValueError("header row is missing")
    header = table[config.header_row - 1]
    if len(set(header)) != len(header) or any(not col for col in header):
        raise ValueError("headers must be unique and nonempty")
    required = [config.amount_column, config.category_column, config.transaction_column]
    if config.row_id_column is not None:
        required.append(config.row_id_column)
    if any(column not in header for column in required):
        raise ValueError("configured column is missing")
    end = config.last_data_row if config.last_data_row is not None else len(table)
    if end < config.header_row or end > len(table):
        raise ValueError("invalid last_data_row")
    rows = []
    for number, cells in enumerate(table[config.header_row : end], config.header_row + 1):
        if len(cells) != len(header):
            raise ValueError(f"row {number}: column count differs from header")
        values = dict(zip(header, cells, strict=True))
        amount = values[config.amount_column]
        if re.fullmatch(r"[+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?", amount) is None:
            raise ValueError(f"row {number}: expected exact decimal amount text")
        rows.append(
            RevenueInputRow(
                row_id=values[config.row_id_column] if config.row_id_column else f"row-{number}",
                amount=Decimal(amount),
                raw_amount=amount,
                raw_category_label=values[config.category_column] or None,
                raw_transaction_label=values[config.transaction_column] or None,
                lineage=Lineage(
                    source_file=path.name,
                    sheet=sheet,
                    row_number=number,
                    raw_hash=raw_hash(cells),
                    run_id=config.run_id,
                    mapping_version=config.adapter_version,
                ),
            )
        )
    return StatementPackage(
        agency_id=config.agency_id,
        carrier=config.carrier,
        period=config.period,
        statement_id=config.statement_id,
        revision_id=config.revision_id,
        content_sha256=hashlib.sha256(data).hexdigest(),
        expected_row_count=len(rows),
        rows=tuple(rows),
        control_total=config.control_total,
    )
