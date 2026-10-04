"""Carrier commission statements: one XLSX per carrier, all three periods on one sheet.

Quirks: row 1 is a merged title, row 2 a free-text subtitle, row 3 the header, data from row
4, and a trailing total row (column A "Total", column B the data row count, the amount column
the sum). Each carrier family has its own column names and order. Money is two-decimal text.
Some layouts store dates as real date cells, others as text.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from synth_agency_data.world import CARRIERS, World
from synth_agency_data.writers.common import Location, fmt_date, save_xlsx, text

SHEET = "Statement"
HEADER_ROW = 3
# Layout: (header, canonical field) pairs in column order, plus the date style (None = cell).
LAYOUTS: dict[str, tuple[tuple[tuple[str, str], ...], str | None]] = {
    "a": (
        (("Line", "line_no"), ("Period", "statement_period"), ("Member ID", "carrier_member_id"),
         ("Member Name", "member_name"), ("DOB", "member_dob"), ("Policy", "policy_ref"),
         ("Writing Agent", "agent_npn"), ("Type", "commission_type"), ("Paid", "paid_date"),
         ("Commission", "amount")),
        None,
    ),
    "b": (
        (("Stmt Period", "statement_period"), ("Seq #", "line_no"),
         ("Subscriber Name", "member_name"), ("Subscriber ID", "carrier_member_id"),
         ("Birth Date", "member_dob"), ("Agent NPN", "agent_npn"),
         ("Agency Policy Ref", "policy_ref"), ("Txn Type", "commission_type"),
         ("Pay Date", "paid_date"), ("Amount Paid", "amount")),
        "mdy",
    ),
    "c": (
        (("Statement Month", "statement_period"), ("Line No", "line_no"),
         ("Member #", "carrier_member_id"), ("Insured", "member_name"),
         ("Date of Birth", "member_dob"), ("Producer NPN", "agent_npn"), ("Ref", "policy_ref"),
         ("Comm Type", "commission_type"), ("Payment Date", "paid_date"), ("Comp $", "amount")),
        "iso",
    ),
}  # fmt: skip
CARRIER_LAYOUT = {
    "Northwind Health": "a",
    "Cardinal Mutual": "a",
    "Bluepeak": "b",
    "Summit Health Plans": "b",
    "Harborline": "c",
    "Meridian Care": "c",
}


def file_name(carrier: str) -> str:
    return f"commissions_{carrier.lower().replace(' ', '_')}.xlsx"


def _cell(value: Any, style: str | None) -> Any:
    if isinstance(value, date):
        return value if style is None else fmt_date(value, style)
    return text(value)


def write_commissions(world: World, out: Path) -> tuple[dict[str, int], dict[str, Location]]:
    """Returns data row counts by file, and where each line landed (keyed carrier|period|no)."""
    counts: dict[str, int] = {}
    where: dict[str, Location] = {}
    for carrier in CARRIERS:
        columns, style = LAYOUTS[CARRIER_LAYOUT[carrier]]
        lines = [ln for ln in world.tables["commission_lines"] if ln["carrier"] == carrier]
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append([f"{carrier} Commission Statement"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        ws.append(["Agency A, statement periods 2026-06 to 2026-08"])
        ws.append([header for header, _ in columns])
        for ln in lines:
            ws.append([_cell(ln[field], style) for _, field in columns])
            key = f"{carrier}|{ln['statement_period']}|{ln['line_no']}"
            where[key] = (file_name(carrier), SHEET, ws.max_row)
        total = [""] * len(columns)
        total[0], total[1] = "Total", str(len(lines))
        total[-1] = text(sum((ln["amount"] for ln in lines), Decimal("0")))
        ws.append(total)
        save_xlsx(wb, out / file_name(carrier))
        counts[file_name(carrier)] = len(lines)
    return counts, where
