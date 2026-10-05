"""Agent roster (a hand-kept workbook): an "Agents" sheet and an "RTS" sheet, header on row 1.

Quirks: free-form headers, license states as one comma list ("FL, GA, TX"), Y/N flags, and
RTS dates as real date cells. With add_ssn, the Agents sheet gets an "SSN" column (SPEC
example 6). Its values all start with area 000, which the Social Security Administration
has never issued, so none can belong to a real person.
"""

import random
from pathlib import Path

from openpyxl import Workbook

from synth_agency_data.world import World
from synth_agency_data.writers.common import save_xlsx, text

FILE = "agent_roster.xlsx"
AGENT_HEADERS = ("Agent Name", "NPN #", "E-mail", "Licensed States", "Upline NPN", "Status")
RTS_HEADERS = (
    "NPN", "Carrier", "State", "Plan Yr", "Line", "Appointed?", "Certified?", "RTS Start",
    "RTS End",
)  # fmt: skip


def fake_ssns(seed: int, n: int) -> list[str]:
    """Distinct, never-issued values: area 000 is invalid by SSA rules."""
    rng = random.Random(f"ssn-{seed}")
    return [f"000-{v // 10000 + 10}-{v % 10000:04d}" for v in rng.sample(range(900_000), n)]


def write_roster(world: World, out: Path, add_ssn: bool) -> tuple[dict[str, int], list[str]]:
    """Returns data row counts by sheet and the SSN values written (empty without add_ssn)."""
    agents, rts = world.tables["agents"], world.tables["rts"]
    ssns = fake_ssns(world.seed, len(agents)) if add_ssn else []
    wb = Workbook()
    ws = wb.active
    ws.title = "Agents"
    ws.append(AGENT_HEADERS + (("SSN",) if add_ssn else ()))
    for i, a in enumerate(agents):
        row = [f"{a['first_name']} {a['last_name']}", a["npn"], text(a["email"])]
        row += [", ".join(a["license_states"]), text(a["upline_npn"]), text(a["status"]).title()]
        ws.append(row + ssns[i : i + 1])
    sheet = wb.create_sheet("RTS")
    sheet.append(RTS_HEADERS)
    for r in rts:
        flags = ["Y" if r[k] else "N" for k in ("appointed", "certified")]
        sheet.append(
            [r["npn"], r["carrier"], r["state"], str(r["plan_year"]), text(r["line_of_business"])]
            + flags
            + [r["effective_date"], r["end_date"]]
        )
    save_xlsx(wb, out / FILE)
    return {"Agents": len(agents), "RTS": len(rts)}, ssns
